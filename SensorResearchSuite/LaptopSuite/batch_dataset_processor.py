#!/usr/bin/env python3
"""
SensorMax Batch Dataset Processor & Comparator

Purpose
-------
Process a directory of SensorMax CSV exports and generate one Excel
workbook containing comparable engineering summaries.

AndroidApp is NOT modified.

Supported current Android CSV formats
--------------------------------------

RAW:

    Timestamp_ms
    Machine_ID
    Point
    Sensor_Type
    Sensor_Name
    Val_0
    Val_1
    Val_2
    Val_3
    Val_4
    Val_5

ANALYSIS:

    Timestamp_ms
    Machine_ID
    Point
    Sample_Rate_Hz
    RMS_X_ms2
    RMS_Y_ms2
    RMS_Z_ms2
    Overall_RMS_Accel_ms2
    Overall_RMS_Velocity_mms
    Overall_RMS_Displacement_um
    Dominant_Axis
    Dominant_Freq_Hz
    Dominant_Accel_Amplitude_ms2
    Envelope_Peak_Freq_Hz
    ISO20816_Zone
    Bearing_Match
    RPM_Input
    Impact_Event
    Snapshot_Triggered

Legacy compatibility
--------------------
Older positional SensorMax CSV files are accepted when possible.

Output workbook
---------------
1. Batch Trial Comparison
2. Analysis Windows
3. Sensor Summary

Engineering intent
------------------
This tool summarizes measured data. It does not claim machine diagnosis.

The previous activity-classification / step-counting workflow is removed
because it does not belong in the SensorMax rotating-equipment vibration
pipeline.
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any


# ============================================================
# Optional dependency check
# ============================================================

try:

    import numpy as np

except ImportError:

    print(
        "ERROR: numpy is required."
    )

    print(
        "Install with:"
    )

    print(
        "    python -m pip install numpy"
    )

    sys.exit(1)


try:

    import openpyxl

    from openpyxl import Workbook

    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )

    from openpyxl.utils import get_column_letter

except ImportError:

    print(
        "ERROR: openpyxl is required."
    )

    print(
        "Install with:"
    )

    print(
        "    python -m pip install openpyxl"
    )

    sys.exit(1)


# ============================================================
# Constants
# ============================================================

CURRENT_RAW_COLUMNS = {
    "Timestamp_ms",
    "Machine_ID",
    "Point",
    "Sensor_Type",
    "Sensor_Name",
    "Val_0",
    "Val_1",
    "Val_2",
}

CURRENT_ANALYSIS_COLUMNS = {
    "Timestamp_ms",
    "Machine_ID",
    "Point",
    "Sample_Rate_Hz",
    "RMS_X_ms2",
    "RMS_Y_ms2",
    "RMS_Z_ms2",
    "Overall_RMS_Accel_ms2",
    "Overall_RMS_Velocity_mms",
    "Overall_RMS_Displacement_um",
    "Dominant_Axis",
    "Dominant_Freq_Hz",
    "Envelope_Peak_Freq_Hz",
    "ISO20816_Zone",
    "Bearing_Match",
    "RPM_Input",
}

SENSOR_ACCELEROMETER = 1

MAX_DISPLAY_RECORDS = 50000


# ============================================================
# Excel styling
# ============================================================

HEADER_FILL = PatternFill(
    start_color="0F172A",
    end_color="0F172A",
    fill_type="solid",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FONT = Font(
    size=16,
    bold=True,
    color="0F172A",
)

SUBTITLE_FONT = Font(
    size=10,
    italic=True,
    color="64748B",
)

GOOD_FILL = PatternFill(
    start_color="DCFCE7",
    end_color="DCFCE7",
    fill_type="solid",
)

WARN_FILL = PatternFill(
    start_color="FEF3C7",
    end_color="FEF3C7",
    fill_type="solid",
)

FAIL_FILL = PatternFill(
    start_color="FEE2E2",
    end_color="FEE2E2",
    fill_type="solid",
)

THIN_SIDE = Side(
    style="thin",
    color="CBD5E1",
)

THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


# ============================================================
# General helpers
# ============================================================

def safe_text(
    value: Any,
    default: str = "",
) -> str:

    if value is None:
        return default

    text = str(value).strip()

    return text if text else default


def safe_float(
    value: Any,
    default: float = 0.0,
) -> float:

    try:

        number = float(value)

        if not math.isfinite(number):
            return default

        return number

    except (
        TypeError,
        ValueError,
    ):

        return default


def safe_int(
    value: Any,
    default: int = 0,
) -> int:

    try:

        return int(
            float(value)
        )

    except (
        TypeError,
        ValueError,
    ):

        return default


def mean_or_zero(
    values: list[float],
) -> float:

    if not values:
        return 0.0

    return float(
        np.mean(values)
    )


def rms(
    values: list[float],
) -> float:

    if not values:
        return 0.0

    arr = np.asarray(
        values,
        dtype=float,
    )

    return float(
        np.sqrt(
            np.mean(
                arr ** 2
            )
        )
    )


def timestamp_to_seconds(
    timestamp_ms: float,
    first_ms: float,
) -> float:

    return (
        timestamp_ms -
        first_ms
    ) / 1000.0


def detect_schema(
    header: list[str],
) -> str:

    columns = {
        str(value).strip()
        for value in header
    }

    if CURRENT_RAW_COLUMNS.issubset(
        columns
    ):

        return "CURRENT_RAW"

    if CURRENT_ANALYSIS_COLUMNS.issubset(
        columns
    ):

        return "CURRENT_ANALYSIS"

    return "LEGACY_OR_UNKNOWN"


def excel_style_header(
    ws,
    row_number: int,
) -> None:

    for cell in ws[row_number]:

        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = THIN_BORDER


def auto_width(
    ws,
    minimum=11,
    maximum=42,
) -> None:

    for column_cells in ws.columns:

        if not column_cells:
            continue

        column_number = (
            column_cells[0].column
        )

        letter = get_column_letter(
            column_number
        )

        longest = 0

        for cell in column_cells:

            text = str(
                cell.value
                if cell.value is not None
                else ""
            )

            longest = max(
                longest,
                len(text)
            )

        ws.column_dimensions[
            letter
        ].width = max(
            minimum,
            min(
                maximum,
                longest + 2,
            ),
        )


def sanitize_sheet_name(
    name: str,
) -> str:

    name = safe_text(
        name,
        "Unknown"
    )

    for char in (
        "[",
        "]",
        ":",
        "*",
        "?",
        "/",
        "\\",
    ):

        name = name.replace(
            char,
            "_",
        )

    return name[:31]


# ============================================================
# CSV loading
# ============================================================

def read_current_raw_csv(
    csv_path: Path,
) -> list[dict[str, Any]]:

    records = []

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(
            file
        )

        for row in reader:

            if not row:
                continue

            record = {
                "timestamp_ms":
                    safe_float(
                        row.get(
                            "Timestamp_ms"
                        )
                    ),

                "machine_id":
                    safe_text(
                        row.get(
                            "Machine_ID"
                        ),
                        "UNSPECIFIED",
                    ),

                "point":
                    safe_text(
                        row.get(
                            "Point"
                        ),
                        "UNSPECIFIED",
                    ),

                "sensor_type":
                    safe_int(
                        row.get(
                            "Sensor_Type"
                        )
                    ),

                "sensor_name":
                    safe_text(
                        row.get(
                            "Sensor_Name"
                        ),
                        "Unknown Sensor",
                    ),

                "v0":
                    safe_float(
                        row.get(
                            "Val_0"
                        )
                    ),

                "v1":
                    safe_float(
                        row.get(
                            "Val_1"
                        )
                    ),

                "v2":
                    safe_float(
                        row.get(
                            "Val_2"
                        )
                    ),

                "v3":
                    safe_float(
                        row.get(
                            "Val_3"
                        )
                    ),

                "v4":
                    safe_float(
                        row.get(
                            "Val_4"
                        )
                    ),

                "v5":
                    safe_float(
                        row.get(
                            "Val_5"
                        )
                    ),
            }

            records.append(
                record
            )

    return records


def read_current_analysis_csv(
    csv_path: Path,
) -> list[dict[str, Any]]:

    records = []

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(
            file
        )

        for row in reader:

            if not row:
                continue

            records.append(
                {
                    "timestamp_ms":
                        safe_float(
                            row.get(
                                "Timestamp_ms"
                            )
                        ),

                    "machine_id":
                        safe_text(
                            row.get(
                                "Machine_ID"
                            ),
                            "UNSPECIFIED",
                        ),

                    "point":
                        safe_text(
                            row.get(
                                "Point"
                            ),
                            "UNSPECIFIED",
                        ),

                    "sample_rate_hz":
                        safe_float(
                            row.get(
                                "Sample_Rate_Hz"
                            )
                        ),

                    "rms_x":
                        safe_float(
                            row.get(
                                "RMS_X_ms2"
                            )
                        ),

                    "rms_y":
                        safe_float(
                            row.get(
                                "RMS_Y_ms2"
                            )
                        ),

                    "rms_z":
                        safe_float(
                            row.get(
                                "RMS_Z_ms2"
                            )
                        ),

                    "overall_rms_accel":
                        safe_float(
                            row.get(
                                "Overall_RMS_Accel_ms2"
                            )
                        ),

                    "overall_rms_velocity":
                        safe_float(
                            row.get(
                                "Overall_RMS_Velocity_mms"
                            )
                        ),

                    "overall_rms_displacement":
                        safe_float(
                            row.get(
                                "Overall_RMS_Displacement_um"
                            )
                        ),

                    "dominant_axis":
                        safe_text(
                            row.get(
                                "Dominant_Axis"
                            )
                        ),

                    "dominant_freq":
                        safe_float(
                            row.get(
                                "Dominant_Freq_Hz"
                            )
                        ),

                    "dominant_accel_amplitude":
                        safe_float(
                            row.get(
                                "Dominant_Accel_Amplitude_ms2"
                            )
                        ),

                    "envelope_peak_freq":
                        safe_float(
                            row.get(
                                "Envelope_Peak_Freq_Hz"
                            )
                        ),

                    "iso_zone":
                        safe_text(
                            row.get(
                                "ISO20816_Zone"
                            )
                        ),

                    "bearing_match":
                        safe_text(
                            row.get(
                                "Bearing_Match"
                            )
                        ),

                    "rpm":
                        safe_float(
                            row.get(
                                "RPM_Input"
                            )
                        ),

                    "impact_event":
                        safe_text(
                            row.get(
                                "Impact_Event"
                            )
                        ),

                    "snapshot":
                        safe_text(
                            row.get(
                                "Snapshot_Triggered"
                            )
                        ),
                }
            )

    return records


def read_legacy_csv(
    csv_path: Path,
) -> list[dict[str, Any]]:

    records = []

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:

        reader = csv.reader(
            file
        )

        header = next(
            reader,
            None,
        )

        if header is None:
            return records

        for row in reader:

            if len(row) < 6:
                continue

            try:

                sensor_type = int(
                    float(
                        row[1]
                    )
                )

            except (
                ValueError,
                TypeError,
            ):

                continue

            # Historical format:
            #
            # timestamp,
            # sensor_type,
            # sensor_name,
            # x,
            # y,
            # z

            records.append(
                {
                    "timestamp_ms":
                        safe_float(
                            row[0]
                        ),

                    "machine_id":
                        "UNSPECIFIED",

                    "point":
                        "UNSPECIFIED",

                    "sensor_type":
                        sensor_type,

                    "sensor_name":
                        safe_text(
                            row[2]
                            if len(row) > 2
                            else "",
                            "Unknown Sensor",
                        ),

                    "v0":
                        safe_float(
                            row[3]
                        ),

                    "v1":
                        safe_float(
                            row[4]
                        ),

                    "v2":
                        safe_float(
                            row[5]
                        ),

                    "v3": 0.0,
                    "v4": 0.0,
                    "v5": 0.0,
                }
            )

    return records


def read_csv_records(
    csv_path: Path,
) -> tuple[str, list[dict[str, Any]]]:

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline="",
    ) as file:

        reader = csv.reader(
            file
        )

        header = next(
            reader,
            None,
        )

    if not header:

        return (
            "EMPTY",
            [],
        )

    schema = detect_schema(
        header
    )

    if schema == "CURRENT_RAW":

        return (
            schema,
            read_current_raw_csv(
                csv_path
            ),
        )

    if schema == "CURRENT_ANALYSIS":

        return (
            schema,
            read_current_analysis_csv(
                csv_path
            ),
        )

    return (
        schema,
        read_legacy_csv(
            csv_path
        ),
    )


# ============================================================
# Raw-record analysis
# ============================================================

def analyze_raw_records(
    records: list[dict[str, Any]],
) -> dict[str, Any]:

    acceleration_records = [
        record
        for record in records
        if (
            record["sensor_type"] ==
            SENSOR_ACCELEROMETER
        )
    ]

    if not acceleration_records:

        # Some devices may report a different sensor type/name.
        acceleration_records = [
            record
            for record in records
            if "accel" in
            record["sensor_name"].lower()
        ]

    if not acceleration_records:

        return {
            "record_count": len(records),
            "accel_count": 0,
            "duration_s": 0.0,
            "effective_rate_hz": 0.0,
            "rms_x": 0.0,
            "rms_y": 0.0,
            "rms_z": 0.0,
            "dynamic_rms": 0.0,
            "peak_to_peak": 0.0,
            "max_rss": 0.0,
            "machine_ids": "",
            "points": "",
            "sensor_names": "",
        }

    timestamps = [
        record["timestamp_ms"]
        for record in acceleration_records
        if record["timestamp_ms"] > 0
    ]

    xs = [
        record["v0"]
        for record in acceleration_records
    ]

    ys = [
        record["v1"]
        for record in acceleration_records
    ]

    zs = [
        record["v2"]
        for record in acceleration_records
    ]

    x_arr = np.asarray(
        xs,
        dtype=float,
    )

    y_arr = np.asarray(
        ys,
        dtype=float,
    )

    z_arr = np.asarray(
        zs,
        dtype=float,
    )

    rss = np.sqrt(
        x_arr ** 2 +
        y_arr ** 2 +
        z_arr ** 2
    )

    # Dynamic vibration magnitude:
    # remove the mean resultant component rather than using raw
    # gravity-included magnitude directly.
    dynamic_rss = (
        rss -
        np.mean(rss)
    )

    if timestamps:

        first_ts = min(
            timestamps
        )

        last_ts = max(
            timestamps
        )

        duration_s = max(
            0.0,
            (
                last_ts -
                first_ts
            ) / 1000.0,
        )

    else:

        duration_s = 0.0

    effective_rate_hz = (
        len(timestamps) /
        duration_s
        if duration_s > 0
        else 0.0
    )

    machine_ids = sorted(
        {
            record["machine_id"]
            for record in acceleration_records
            if record["machine_id"]
        }
    )

    points = sorted(
        {
            record["point"]
            for record in acceleration_records
            if record["point"]
        }
    )

    sensor_names = sorted(
        {
            record["sensor_name"]
            for record in acceleration_records
            if record["sensor_name"]
        }
    )

    return {
        "record_count":
            len(records),

        "accel_count":
            len(acceleration_records),

        "duration_s":
            duration_s,

        "effective_rate_hz":
            effective_rate_hz,

        "rms_x":
            rms(xs),

        "rms_y":
            rms(ys),

        "rms_z":
            rms(zs),

        "dynamic_rms":
            rms(
                dynamic_rss.tolist()
            ),

        "peak_to_peak":
            float(
                np.max(rss) -
                np.min(rss)
            ),

        "max_rss":
            float(
                np.max(rss)
            ),

        "machine_ids":
            ", ".join(machine_ids),

        "points":
            ", ".join(points),

        "sensor_names":
            ", ".join(sensor_names),
    }


# ============================================================
# Analysis-record summary
# ============================================================

def analyze_analysis_records(
    records: list[dict[str, Any]],
) -> dict[str, Any]:

    if not records:

        return {
            "window_count": 0,
            "duration_s": 0.0,
            "mean_sample_rate_hz": 0.0,
            "mean_rms_velocity": 0.0,
            "max_rms_velocity": 0.0,
            "mean_rms_accel": 0.0,
            "max_rms_accel": 0.0,
            "mean_dominant_freq": 0.0,
            "max_dominant_freq": 0.0,
            "mean_envelope_freq": 0.0,
            "rpm_values": "",
            "iso_zones": "",
            "bearing_matches": "",
            "impact_count": 0,
            "snapshot_count": 0,
        }

    def numeric(
        key: str,
    ) -> list[float]:

        return [
            record[key]
            for record in records
            if math.isfinite(
                record[key]
            )
        ]

    timestamps = [
        record["timestamp_ms"]
        for record in records
        if record["timestamp_ms"] > 0
    ]

    duration_s = (
        (
            max(timestamps) -
            min(timestamps)
        ) / 1000.0
        if len(timestamps) >= 2
        else 0.0
    )

    sample_rates = numeric(
        "sample_rate_hz"
    )

    rms_velocities = numeric(
        "overall_rms_velocity"
    )

    rms_accels = numeric(
        "overall_rms_accel"
    )

    dominant_freqs = numeric(
        "dominant_freq"
    )

    envelope_freqs = numeric(
        "envelope_peak_freq"
    )

    rpm_values = sorted(
        {
            round(
                value,
                2,
            )
            for value in numeric(
                "rpm"
            )
            if value > 0
        }
    )

    iso_zones = sorted(
        {
            record["iso_zone"]
            for record in records
            if record["iso_zone"]
        }
    )

    bearing_matches = sorted(
        {
            record["bearing_match"]
            for record in records
            if record["bearing_match"]
        }
    )

    impact_count = sum(
        1
        for record in records
        if record["impact_event"]
        .strip()
        .lower()
        in {
            "true",
            "1",
            "yes",
        }
    )

    snapshot_count = sum(
        1
        for record in records
        if record["snapshot"]
        .strip()
        .lower()
        in {
            "true",
            "1",
            "yes",
        }
    )

    return {
        "window_count":
            len(records),

        "duration_s":
            duration_s,

        "mean_sample_rate_hz":
            mean_or_zero(
                sample_rates
            ),

        "mean_rms_velocity":
            mean_or_zero(
                rms_velocities
            ),

        "max_rms_velocity":
            max(
                rms_velocities,
                default=0.0,
            ),

        "mean_rms_accel":
            mean_or_zero(
                rms_accels
            ),

        "max_rms_accel":
            max(
                rms_accels,
                default=0.0,
            ),

        "mean_dominant_freq":
            mean_or_zero(
                dominant_freqs
            ),

        "max_dominant_freq":
            max(
                dominant_freqs,
                default=0.0,
            ),

        "mean_envelope_freq":
            mean_or_zero(
                envelope_freqs
            ),

        "rpm_values":
            ", ".join(
                str(value)
                for value in rpm_values
            ),

        "iso_zones":
            ", ".join(
                iso_zones
            ),

        "bearing_matches":
            ", ".join(
                bearing_matches
            ),

        "impact_count":
            impact_count,

        "snapshot_count":
            snapshot_count,
    }


# ============================================================
# Workbook: Batch Trial Comparison
# ============================================================

def create_batch_summary_sheet(
    workbook: Workbook,
    rows: list[dict[str, Any]],
) -> None:

    ws = workbook.active

    ws.title = (
        "Batch Trial Comparison"
    )

    ws["A1"] = (
        "SensorMax Batch Trial Comparison Report"
    )

    ws["A1"].font = TITLE_FONT

    ws["A2"] = (
        "Engineering comparison of imported SensorMax sessions"
    )

    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "File Name",
        "Schema",
        "Machine ID",
        "Point",
        "Records / Windows",
        "Acceleration Records",
        "Duration (s)",
        "Effective Rate (Hz)",
        "Mean Sample Rate (Hz)",
        "RMS X (m/s²)",
        "RMS Y (m/s²)",
        "RMS Z (m/s²)",
        "Dynamic RMS (m/s²)",
        "RMS Velocity (mm/s)",
        "Peak RMS Velocity (mm/s)",
        "Peak-to-Peak RSS (m/s²)",
        "Dominant Frequency (Hz)",
        "Envelope Peak (Hz)",
        "RPM",
        "ISO Zone",
        "Bearing Match",
        "Impact Events",
        "Snapshots",
    ]

    header_row = 4

    for col, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            row=header_row,
            column=col,
            value=header,
        )

    excel_style_header(
        ws,
        header_row,
    )

    row_number = 5

    for item in rows:

        ws.append(
            [
                item["file_name"],
                item["schema"],
                item["machine_id"],
                item["point"],
                item["record_count"],
                item["accel_count"],
                item["duration_s"],
                item["effective_rate_hz"],
                item["mean_sample_rate_hz"],
                item["rms_x"],
                item["rms_y"],
                item["rms_z"],
                item["dynamic_rms"],
                item["mean_rms_velocity"],
                item["max_rms_velocity"],
                item["peak_to_peak"],
                item["mean_dominant_freq"],
                item["mean_envelope_freq"],
                item["rpm"],
                item["iso_zones"],
                item["bearing_matches"],
                item["impact_count"],
                item["snapshot_count"],
            ]
        )

        row_number += 1

    for row in ws.iter_rows(
        min_row=5,
        max_row=max(
            5,
            row_number - 1,
        ),
    ):

        for cell in row:

            cell.border = THIN_BORDER

    for row in ws.iter_rows(
        min_row=5,
        max_row=max(
            5,
            row_number - 1,
        ),
        min_col=7,
        max_col=18,
    ):

        for cell in row:
            cell.number_format = (
                "0.000"
            )

    ws.freeze_panes = (
        "A5"
    )

    if row_number > 5:

        ws.auto_filter.ref = (
            f"A4:W{row_number - 1}"
        )

    auto_width(ws)


# ============================================================
# Workbook: Analysis Windows
# ============================================================

def create_analysis_sheet(
    workbook: Workbook,
    analysis_rows: list[dict[str, Any]],
) -> None:

    ws = workbook.create_sheet(
        "Analysis Windows"
    )

    ws["A1"] = (
        "SensorMax Analysis Windows"
    )

    ws["A1"].font = TITLE_FONT

    ws["A2"] = (
        "Detailed records from Android *_analysis.csv files"
    )

    ws["A2"].font = SUBTITLE_FONT

    headers = [
        "File",
        "Timestamp (ms)",
        "Machine ID",
        "Point",
        "Sample Rate (Hz)",
        "RMS X (m/s²)",
        "RMS Y (m/s²)",
        "RMS Z (m/s²)",
        "Overall RMS Accel (m/s²)",
        "Overall RMS Velocity (mm/s)",
        "Overall RMS Displacement (µm)",
        "Dominant Axis",
        "Dominant Frequency (Hz)",
        "Dominant Accel Amplitude (m/s²)",
        "Envelope Peak (Hz)",
        "ISO Zone",
        "Bearing Match",
        "RPM",
        "Impact",
        "Snapshot",
    ]

    header_row = 4

    for col, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            row=header_row,
            column=col,
            value=header,
        )

    excel_style_header(
        ws,
        header_row,
    )

    row_number = 5

    for record in analysis_rows:

        ws.append(
            [
                record["file_name"],
                record["timestamp_ms"],
                record["machine_id"],
                record["point"],
                record["sample_rate_hz"],
                record["rms_x"],
                record["rms_y"],
                record["rms_z"],
                record["overall_rms_accel"],
                record["overall_rms_velocity"],
                record["overall_rms_displacement"],
                record["dominant_axis"],
                record["dominant_freq"],
                record["dominant_accel_amplitude"],
                record["envelope_peak_freq"],
                record["iso_zone"],
                record["bearing_match"],
                record["rpm"],
                record["impact_event"],
                record["snapshot"],
            ]
        )

        row_number += 1

    if row_number > 5:

        for row in ws.iter_rows(
            min_row=5,
            max_row=row_number - 1,
        ):

            for cell in row:
                cell.border = THIN_BORDER

        ws.auto_filter.ref = (
            f"A4:T{row_number - 1}"
        )

    ws.freeze_panes = (
        "A5"
    )

    auto_width(ws)


# ============================================================
# Workbook: Sensor Summary
# ============================================================

def create_sensor_summary_sheet(
    workbook: Workbook,
    sensor_summary: list[dict[str, Any]],
) -> None:

    ws = workbook.create_sheet(
        "Sensor Summary"
    )

    ws["A1"] = (
        "SensorMax Sensor Summary"
    )

    ws["A1"].font = TITLE_FONT

    headers = [
        "File",
        "Sensor Type",
        "Sensor Name",
        "Record Count",
        "Machine IDs",
        "Measurement Points",
        "Mean Value 0",
        "RMS Value 0",
        "Mean Value 1",
        "RMS Value 1",
        "Mean Value 2",
        "RMS Value 2",
    ]

    header_row = 3

    for col, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            row=header_row,
            column=col,
            value=header,
        )

    excel_style_header(
        ws,
        header_row,
    )

    row_number = 4

    for item in sensor_summary:

        ws.append(
            [
                item["file_name"],
                item["sensor_type"],
                item["sensor_name"],
                item["record_count"],
                item["machine_ids"],
                item["points"],
                item["mean_v0"],
                item["rms_v0"],
                item["mean_v1"],
                item["rms_v1"],
                item["mean_v2"],
                item["rms_v2"],
            ]
        )

        row_number += 1

    if row_number > 4:

        ws.auto_filter.ref = (
            f"A3:L{row_number - 1}"
        )

    ws.freeze_panes = (
        "A4"
    )

    auto_width(ws)


# ============================================================
# Per-file processing
# ============================================================

def process_file(
    csv_path: Path,
) -> tuple[
    dict[str, Any],
    list[dict[str, Any]],
    list[dict[str, Any]],
]:

    schema, records = read_csv_records(
        csv_path
    )

    print(
        f"[FILE] {csv_path.name}"
    )

    print(
        f"       Schema : {schema}"
    )

    print(
        f"       Records: {len(records)}"
    )

    batch_row = {
        "file_name":
            csv_path.name,

        "schema":
            schema,

        "machine_id":
            "UNSPECIFIED",

        "point":
            "UNSPECIFIED",

        "record_count":
            len(records),

        "accel_count":
            0,

        "duration_s":
            0.0,

        "effective_rate_hz":
            0.0,

        "mean_sample_rate_hz":
            0.0,

        "rms_x":
            0.0,

        "rms_y":
            0.0,

        "rms_z":
            0.0,

        "dynamic_rms":
            0.0,

        "mean_rms_velocity":
            0.0,

        "max_rms_velocity":
            0.0,

        "peak_to_peak":
            0.0,

        "mean_dominant_freq":
            0.0,

        "mean_envelope_freq":
            0.0,

        "rpm":
            "",

        "iso_zones":
            "",

        "bearing_matches":
            "",

        "impact_count":
            0,

        "snapshot_count":
            0,
    }

    analysis_rows = []

    sensor_rows = []

    # --------------------------------------------------------
    # Current or legacy RAW
    # --------------------------------------------------------

    if (
        schema ==
        "CURRENT_RAW"
        or schema ==
        "LEGACY_OR_UNKNOWN"
    ):

        raw_summary = (
            analyze_raw_records(
                records
            )
        )

        machine_ids = sorted(
            {
                record["machine_id"]
                for record in records
                if record["machine_id"]
            }
        )

        points = sorted(
            {
                record["point"]
                for record in records
                if record["point"]
            }
        )

        if machine_ids:

            batch_row["machine_id"] = (
                ", ".join(
                    machine_ids
                )
            )

        if points:

            batch_row["point"] = (
                ", ".join(
                    points
                )
            )

        batch_row["record_count"] = (
            raw_summary[
                "record_count"
            ]
        )

        batch_row["accel_count"] = (
            raw_summary[
                "accel_count"
            ]
        )

        batch_row["duration_s"] = (
            raw_summary[
                "duration_s"
            ]
        )

        batch_row["effective_rate_hz"] = (
            raw_summary[
                "effective_rate_hz"
            ]
        )

        batch_row["rms_x"] = (
            raw_summary[
                "rms_x"
            ]
        )

        batch_row["rms_y"] = (
            raw_summary[
                "rms_y"
            ]
        )

        batch_row["rms_z"] = (
            raw_summary[
                "rms_z"
            ]
        )

        batch_row["dynamic_rms"] = (
            raw_summary[
                "dynamic_rms"
            ]
        )

        batch_row["peak_to_peak"] = (
            raw_summary[
                "peak_to_peak"
            ]
        )

        # Build per-sensor summary.
        grouped = defaultdict(
            list
        )

        for record in records:

            grouped[
                (
                    record[
                        "sensor_type"
                    ],
                    record[
                        "sensor_name"
                    ],
                )
            ].append(
                record
            )

        for (
            (sensor_type, sensor_name),
            group,
        ) in sorted(
            grouped.items(),
            key=lambda item:
                (
                    item[0][0],
                    item[0][1],
                ),
        ):

            v0 = [
                record["v0"]
                for record in group
            ]

            v1 = [
                record["v1"]
                for record in group
            ]

            v2 = [
                record["v2"]
                for record in group
            ]

            machine_values = sorted(
                {
                    record["machine_id"]
                    for record in group
                    if record["machine_id"]
                }
            )

            point_values = sorted(
                {
                    record["point"]
                    for record in group
                    if record["point"]
                }
            )

            sensor_rows.append(
                {
                    "file_name":
                        csv_path.name,

                    "sensor_type":
                        sensor_type,

                    "sensor_name":
                        sensor_name,

                    "record_count":
                        len(group),

                    "machine_ids":
                        ", ".join(
                            machine_values
                        ),

                    "points":
                        ", ".join(
                            point_values
                        ),

                    "mean_v0":
                        mean_or_zero(v0),

                    "rms_v0":
                        rms(v0),

                    "mean_v1":
                        mean_or_zero(v1),

                    "rms_v1":
                        rms(v1),

                    "mean_v2":
                        mean_or_zero(v2),

                    "rms_v2":
                        rms(v2),
                }
            )

    # --------------------------------------------------------
    # Current ANALYSIS
    # --------------------------------------------------------

    elif schema == "CURRENT_ANALYSIS":

        analysis_summary = (
            analyze_analysis_records(
                records
            )
        )

        machine_ids = sorted(
            {
                record["machine_id"]
                for record in records
                if record["machine_id"]
            }
        )

        points = sorted(
            {
                record["point"]
                for record in records
                if record["point"]
            }
        )

        batch_row["machine_id"] = (
            ", ".join(
                machine_ids
            )
        )

        batch_row["point"] = (
            ", ".join(
                points
            )
        )

        batch_row["record_count"] = (
            analysis_summary[
                "window_count"
            ]
        )

        batch_row["duration_s"] = (
            analysis_summary[
                "duration_s"
            ]
        )

        batch_row[
            "mean_sample_rate_hz"
        ] = (
            analysis_summary[
                "mean_sample_rate_hz"
            ]
        )

        batch_row[
            "mean_rms_velocity"
        ] = (
            analysis_summary[
                "mean_rms_velocity"
            ]
        )

        batch_row[
            "max_rms_velocity"
        ] = (
            analysis_summary[
                "max_rms_velocity"
            ]
        )

        batch_row[
            "rms_x"
        ] = mean_or_zero(
            [
                record["rms_x"]
                for record in records
            ]
        )

        batch_row[
            "rms_y"
        ] = mean_or_zero(
            [
                record["rms_y"]
                for record in records
            ]
        )

        batch_row[
            "rms_z"
        ] = mean_or_zero(
            [
                record["rms_z"]
                for record in records
            ]
        )

        batch_row[
            "mean_dominant_freq"
        ] = (
            analysis_summary[
                "mean_dominant_freq"
            ]
        )

        batch_row[
            "mean_envelope_freq"
        ] = (
            analysis_summary[
                "mean_envelope_freq"
            ]
        )

        batch_row["rpm"] = (
            analysis_summary[
                "rpm_values"
            ]
        )

        batch_row["iso_zones"] = (
            analysis_summary[
                "iso_zones"
            ]
        )

        batch_row[
            "bearing_matches"
        ] = (
            analysis_summary[
                "bearing_matches"
            ]
        )

        batch_row[
            "impact_count"
        ] = (
            analysis_summary[
                "impact_count"
            ]
        )

        batch_row[
            "snapshot_count"
        ] = (
            analysis_summary[
                "snapshot_count"
            ]
        )

        for record in records:

            item = dict(
                record
            )

            item["file_name"] = (
                csv_path.name
            )

            analysis_rows.append(
                item
            )

    return (
        batch_row,
        analysis_rows,
        sensor_rows,
    )


# ============================================================
# Main batch processor
# ============================================================

def process_batch(
    directory_path: str,
    output_xlsx: str = (
        "master_batch_comparison.xlsx"
    ),
) -> None:

    directory = (
        Path(
            directory_path
        )
        .expanduser()
        .resolve()
    )

    output_path = (
        Path(
            output_xlsx
        )
        .expanduser()
        .resolve()
    )

    if not directory.exists():

        print(
            f"ERROR: Directory does not exist: "
            f"{directory}"
        )

        return

    if not directory.is_dir():

        print(
            f"ERROR: Not a directory: "
            f"{directory}"
        )

        return

    csv_files = sorted(
        directory.glob(
            "*.csv"
        )
    )

    if not csv_files:

        print(
            f"No CSV files found in: "
            f"{directory}"
        )

        return

    print(
        "============================================================"
    )

    print(
        "  SENSORMAX BATCH DATASET PROCESSOR"
    )

    print(
        "============================================================"
    )

    print(
        f"Input directory : {directory}"
    )

    print(
        f"CSV files       : {len(csv_files)}"
    )

    print(
        f"Output workbook : {output_path}"
    )

    print(
        "============================================================"
    )

    batch_rows = []

    analysis_rows = []

    sensor_rows = []

    for csv_path in csv_files:

        try:

            (
                batch_row,
                file_analysis_rows,
                file_sensor_rows,
            ) = process_file(
                csv_path
            )

            batch_rows.append(
                batch_row
            )

            analysis_rows.extend(
                file_analysis_rows
            )

            sensor_rows.extend(
                file_sensor_rows
            )

        except Exception as exc:

            print(
                f"[ERROR] Failed: "
                f"{csv_path.name}"
            )

            print(
                f"        {type(exc).__name__}: "
                f"{exc}"
            )

    # --------------------------------------------------------
    # Build workbook
    # --------------------------------------------------------

    workbook = Workbook()

    create_batch_summary_sheet(
        workbook,
        batch_rows,
    )

    create_analysis_sheet(
        workbook,
        analysis_rows,
    )

    create_sensor_summary_sheet(
        workbook,
        sensor_rows,
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(
        output_path
    )

    print(
        "============================================================"
    )

    print(
        "[SUCCESS] Batch workbook generated"
    )

    print(
        f"          {output_path}"
    )

    print(
        f"          Trials/files     : {len(batch_rows)}"
    )

    print(
        f"          Analysis windows : {len(analysis_rows)}"
    )

    print(
        f"          Sensor groups    : {len(sensor_rows)}"
    )

    print(
        "============================================================"
    )


# ============================================================
# CLI
# ============================================================

def build_parser() -> argparse.ArgumentParser:

    parser = argparse.ArgumentParser(
        description=(
            "Process a directory of SensorMax CSV "
            "datasets into one comparison workbook."
        )
    )

    parser.add_argument(
        "directory",
        help=(
            "Directory containing SensorMax CSV files"
        ),
    )

    parser.add_argument(
        "--output",
        default=(
            "master_batch_comparison.xlsx"
        ),
        help=(
            "Output Excel workbook path"
        ),
    )

    return parser


def main() -> int:

    parser = build_parser()

    args = parser.parse_args()

    try:

        process_batch(
            args.directory,
            args.output,
        )

        return 0

    except KeyboardInterrupt:

        print(
            "\n[INFO] Batch processing cancelled."
        )

        return 130

    except Exception as exc:

        print(
            "[ERROR] Batch processing failed: "
            f"{type(exc).__name__}: {exc}"
        )

        return 1


if __name__ == "__main__":

    raise SystemExit(
        main()
    )