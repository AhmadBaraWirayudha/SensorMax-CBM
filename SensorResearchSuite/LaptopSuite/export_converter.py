#!/usr/bin/env python3
"""
SensorMax Sensor Data Exporter

Purpose
-------
Convert SensorMax raw CSV files into a modern multi-sheet Excel
workbook grouped by Android sensor type.

Current Android raw CSV schema
------------------------------

Schema v2 adds distinct sensor-clock time, source/channel provenance and units:

    Schema_Version, Sensor_Timestamp_ns, Timestamp_ms, Machine_ID, Point,
    Source_Kind, Source_ID, Channel_ID, Sensor_Type, Sensor_Name, Unit,
    Val_0, Val_1, Val_2, Val_3, Val_4, Val_5

The previous 11-column header remains supported as schema v1.

The Android application is treated as the source of truth.

This converter:

    CSV
      |
      +--> metadata
      |
      +--> one worksheet per sensor type
      |
      +--> normalized sensor values
      |
      +--> relative time
      |
      +--> summary statistics
      |
      +--> engineering-friendly formatting

Backward compatibility
----------------------
Older CSV files using positional columns are accepted when possible.

AndroidApp is NOT modified by this script.
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


# ============================================================
# Optional dependency check
# ============================================================

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print(
        "ERROR: openpyxl is required."
    )
    print(
        "Install it with:"
    )
    print(
        "    python -m pip install openpyxl"
    )
    sys.exit(1)


# ============================================================
# Sensor metadata
# ============================================================

SENSOR_NAMES = {
    1: "Accelerometer",
    2: "Magnetometer",
    4: "Gyroscope",
    5: "Light Sensor",
    8: "Proximity Sensor",
    9: "Gravity Sensor",
    10: "Linear Acceleration",
    11: "Rotation Vector",
}


def sensor_name(sensor_type: int, reported_name: str = "") -> str:
    """
    Prefer the actual Android sensor name when available.
    Fall back to the Android sensor-type mapping.
    """

    reported_name = str(
        reported_name or ""
    ).strip()

    if reported_name:
        return reported_name

    return SENSOR_NAMES.get(
        sensor_type,
        f"Sensor Type {sensor_type}",
    )


# ============================================================
# Styling
# ============================================================

NAVY = "1E3A8A"
BLUE = "2563EB"
LIGHT_BLUE = "DBEAFE"
DARK = "111827"
GRAY = "6B7280"
LIGHT_GRAY = "E5E7EB"
WHITE = "FFFFFF"
GREEN = "166534"
LIGHT_GREEN = "DCFCE7"


HEADER_FILL = PatternFill(
    start_color=NAVY,
    end_color=NAVY,
    fill_type="solid",
)

SUBHEADER_FILL = PatternFill(
    start_color=LIGHT_BLUE,
    end_color=LIGHT_BLUE,
    fill_type="solid",
)

GOOD_FILL = PatternFill(
    start_color=LIGHT_GREEN,
    end_color=LIGHT_GREEN,
    fill_type="solid",
)

HEADER_FONT = Font(
    color=WHITE,
    bold=True,
)

TITLE_FONT = Font(
    size=15,
    bold=True,
    color=NAVY,
)

SUBTITLE_FONT = Font(
    size=10,
    italic=True,
    color=GRAY,
)

BOLD_FONT = Font(
    bold=True,
    color=DARK,
)

THIN_SIDE = Side(
    style="thin",
    color=LIGHT_GRAY,
)

THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


# ============================================================
# Utility functions
# ============================================================

def safe_float(value: Any, default: float | None = 0.0) -> float | None:
    """Parse a finite CSV number, preserving missing/invalid values when default=None."""
    if value is None:
        return default
    text = str(value).strip()
    if not text:
        return default
    try:
        number = float(text)
    except (TypeError, ValueError, OverflowError):
        return default
    return number if math.isfinite(number) else default


def safe_int(value: Any, default: int | None = 0) -> int | None:
    """Parse an exact integer without routing large timestamps through float."""
    if value is None:
        return default
    text = str(value).strip()
    if not text:
        return default
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError, TypeError):
        return default
    if not number.is_finite() or number != number.to_integral_value():
        return default
    try:
        return int(number)
    except (OverflowError, ValueError):
        return default


def clean_text(value: Any) -> str:
    """
    Convert a value to clean worksheet text.
    """

    if value is None:
        return ""

    return str(value).strip()


def safe_sheet_name(name: str) -> str:
    """
    Excel worksheet names:
      - max 31 characters
      - cannot contain: []:*?/\\
    """

    name = clean_text(name)

    if not name:
        name = "Unknown Sensor"

    invalid = (
        "[",
        "]",
        ":",
        "*",
        "?",
        "/",
        "\\",
    )

    for char in invalid:
        name = name.replace(
            char,
            "_",
        )

    return name[:31]


def unique_sheet_name(
    workbook: Workbook,
    proposed: str,
) -> str:
    """
    Guarantee a unique Excel worksheet name.
    """

    base = safe_sheet_name(
        proposed
    )

    existing = {
        str(name).lower()
        for name in workbook.sheetnames
    }

    if base.lower() not in existing:
        return base

    index = 2

    while True:

        suffix = f"_{index}"

        candidate = (
            base[: 31 - len(suffix)] +
            suffix
        )

        if candidate.lower() not in existing:
            return candidate

        index += 1


def parse_timestamp(value: Any) -> float:
    """
    SensorMax timestamps are epoch milliseconds.
    """

    return safe_float(
        value,
        0.0,
    )


def format_timestamp_ms(timestamp_ms: float) -> str:
    """
    Human-readable timestamp for metadata.
    """

    if timestamp_ms <= 0:
        return "--"

    try:

        dt = datetime.fromtimestamp(
            timestamp_ms / 1000.0
        )

        return dt.strftime(
            "%Y-%m-%d %H:%M:%S"
        )

    except (OverflowError, OSError, ValueError):

        return "--"


# ============================================================
# CSV parsing
# ============================================================

def detect_current_android_schema(
    fieldnames: list[str] | None,
) -> bool:

    if not fieldnames:
        return False

    required = {
        "Timestamp_ms",
        "Machine_ID",
        "Point",
        "Sensor_Type",
        "Sensor_Name",
        "Val_0",
        "Val_1",
        "Val_2",
    }

    return required.issubset(
        set(fieldnames)
    )


def normalize_current_row(
    row: dict[str, Any],
) -> dict[str, Any]:
    """
    Convert a current Android CSV row into a stable internal record.
    """

    sensor_type = safe_int(
        row.get("Sensor_Type"),
        -1,
    )

    reported_name = clean_text(
        row.get("Sensor_Name")
    )

    return {
        "schema_version": safe_int(row.get("Schema_Version"), 1),
        "sensor_timestamp_ns": safe_int(row.get("Sensor_Timestamp_ns"), None),
        "timestamp_ms": parse_timestamp(row.get("Timestamp_ms")),
        "machine_id": clean_text(row.get("Machine_ID")) or "UNSPECIFIED",
        "point": clean_text(row.get("Point")) or "UNSPECIFIED",
        "source_kind": clean_text(row.get("Source_Kind")),
        "source_id": clean_text(row.get("Source_ID")),
        "channel_id": clean_text(row.get("Channel_ID")),
        "sensor_type": sensor_type,
        "sensor_name": sensor_name(sensor_type, reported_name),
        "unit": clean_text(row.get("Unit")),
        # An unavailable/invalid numeric sample must remain unavailable; do not
        # translate it into a valid-looking zero during normalization.
        "v0": safe_float(row.get("Val_0"), None),
        "v1": safe_float(row.get("Val_1"), None),
        "v2": safe_float(row.get("Val_2"), None),
        "v3": safe_float(row.get("Val_3"), None),
        "v4": safe_float(row.get("Val_4"), None),
        "v5": safe_float(row.get("Val_5"), None),
    }


def normalize_legacy_row(
    header: list[str],
    row: list[str],
) -> dict[str, Any] | None:
    """
    Best-effort parser for the older positional format.

    Historical files commonly looked approximately like:

        Timestamp_ms,
        Sensor_Type,
        Sensor_Name,
        Value_X,
        Value_Y,
        Value_Z

    This path exists only for backwards compatibility.
    """

    if len(row) < 2:
        return None

    timestamp_ms = safe_float(
        row[0],
        0.0,
    )

    sensor_type = safe_int(
        row[1],
        0,
    )

    name = (
        row[2].strip('"')
        if len(row) > 2
        else ""
    )

    v0 = (
        safe_float(row[3])
        if len(row) > 3
        else 0.0
    )

    v1 = (
        safe_float(row[4])
        if len(row) > 4
        else 0.0
    )

    v2 = (
        safe_float(row[5])
        if len(row) > 5
        else 0.0
    )

    return {
        "timestamp_ms": timestamp_ms,
        "machine_id": "UNSPECIFIED",
        "point": "UNSPECIFIED",
        "sensor_type": sensor_type,
        "sensor_name": sensor_name(
            sensor_type,
            name,
        ),
        "v0": v0,
        "v1": v1,
        "v2": v2,
        "v3": 0.0,
        "v4": 0.0,
        "v5": 0.0,
    }


def read_csv(
    csv_path: Path,
) -> tuple[list[dict[str, Any]], bool]:
    """
    Read and normalize a SensorMax CSV.

    Returns:
        records,
        used_current_android_schema
    """

    if not csv_path.exists():

        raise FileNotFoundError(
            f"CSV file not found: {csv_path}"
        )

    records: list[dict[str, Any]] = []

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.reader(f)

        header = next(
            reader,
            None,
        )

        if not header:
            return [], False

    normalized_header = [
        str(value).strip()
        for value in header
    ]

    if detect_current_android_schema(
        normalized_header
    ):

        with csv_path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as f:

            dict_reader = csv.DictReader(f)

            for row in dict_reader:

                if not row:
                    continue

                records.append(
                    normalize_current_row(
                        row
                    )
                )

        return records, True

    # --------------------------------------------------------
    # Legacy parser
    # --------------------------------------------------------

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as f:

        reader = csv.reader(f)

        next(
            reader,
            None,
        )

        for row in reader:

            normalized = (
                normalize_legacy_row(
                    normalized_header,
                    row,
                )
            )

            if normalized is not None:
                records.append(
                    normalized
                )

    return records, False


# ============================================================
# Statistics
# ============================================================

def statistics(
    records: list[dict[str, Any]],
) -> dict[str, float | int | None]:
    """Summarize only finite, available values for each channel."""
    result: dict[str, float | int | None] = {"count": len(records)}

    def values(key: str) -> list[float]:
        output = []
        for record in records:
            value = record.get(key)
            if value is None:
                continue
            try:
                number = float(value)
            except (TypeError, ValueError, OverflowError):
                continue
            if math.isfinite(number):
                output.append(number)
        return output

    def rms(data: list[float]) -> float | None:
        if not data:
            return None
        return math.sqrt(sum(value * value for value in data) / len(data))

    for key in ("v0", "v1", "v2"):
        data = values(key)
        result[f"min_{key}"] = min(data) if data else None
        result[f"max_{key}"] = max(data) if data else None
        result[f"mean_{key}"] = sum(data) / len(data) if data else None
        result[f"rms_{key}"] = rms(data)

    return result


# ============================================================
# Workbook helpers
# ============================================================

def style_header_row(
    ws,
    row_number: int,
) -> None:

    for cell in ws[row_number]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = THIN_BORDER


def force_formula_like_input_to_text(wb: Workbook) -> None:
    """Store formula-like CSV text as a literal string in XLSX output.

    The exporter does not create intentional formulas. Any formula-typed cell
    therefore came from imported text and must be stored as text before saving.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    cell.data_type = "s"


def auto_width(
    ws,
    minimum: int = 12,
    maximum: int = 45,
) -> None:

    for column_cells in ws.columns:

        if not column_cells:
            continue

        column_index = (
            column_cells[0].column
        )

        letter = get_column_letter(
            column_index
        )

        longest = 0

        for cell in column_cells:

            try:
                value_length = len(
                    str(
                        cell.value
                        if cell.value is not None
                        else ""
                    )
                )
            except Exception:
                value_length = 0

            longest = max(
                longest,
                value_length,
            )

        ws.column_dimensions[
            letter
        ].width = max(
            minimum,
            min(
                maximum,
                longest + 2,
            ),
        )


def freeze_header(
    ws,
    row: int,
) -> None:

    ws.freeze_panes = (
        f"A{row + 1}"
    )


# ============================================================
# Metadata sheet
# ============================================================

def create_metadata_sheet(
    wb: Workbook,
    csv_path: Path,
    records: list[dict[str, Any]],
    current_schema: bool,
) -> None:

    ws = wb.create_sheet(
        title="Session Metadata"
    )

    ws["A1"] = "SensorMax Dataset"
    ws["A1"].font = TITLE_FONT

    ws["A2"] = (
        "Generated from Android SensorMax raw CSV"
    )
    ws["A2"].font = SUBTITLE_FONT

    rows = [
        ("Source file", csv_path.name),
        (
            "Input schema",
            "Current Android SensorMax"
            if current_schema
            else "Legacy compatibility",
        ),
        (
            "Records",
            len(records),
        ),
    ]

    if records:

        timestamps = [
            r["timestamp_ms"]
            for r in records
            if r["timestamp_ms"] > 0
        ]

        machine_ids = sorted(
            {
                r["machine_id"]
                for r in records
            }
        )

        points = sorted(
            {
                r["point"]
                for r in records
            }
        )

        sensor_types = sorted(
            {
                r["sensor_type"]
                for r in records
            }
        )

        rows.extend(
            [
                (
                    "Machine IDs",
                    ", ".join(
                        machine_ids
                    ),
                ),
                (
                    "Measurement points",
                    ", ".join(points),
                ),
                (
                    "First timestamp",
                    format_timestamp_ms(
                        min(timestamps)
                    )
                    if timestamps
                    else "--",
                ),
                (
                    "Last timestamp",
                    format_timestamp_ms(
                        max(timestamps)
                    )
                    if timestamps
                    else "--",
                ),
                (
                    "Sensor types",
                    ", ".join(
                        str(x)
                        for x in sensor_types
                    ),
                ),
            ]
        )

    row_number = 4

    for label, value in rows:

        ws.cell(
            row=row_number,
            column=1,
            value=label,
        )

        ws.cell(
            row=row_number,
            column=1,
        ).font = BOLD_FONT

        ws.cell(
            row=row_number,
            column=1,
        ).fill = SUBHEADER_FILL

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

        ws.cell(
            row=row_number,
            column=1,
        ).border = THIN_BORDER

        ws.cell(
            row=row_number,
            column=2,
        ).border = THIN_BORDER

        row_number += 1

    row_number += 1

    ws.cell(
        row=row_number,
        column=1,
        value="Sensor Summary",
    )

    ws.cell(
        row=row_number,
        column=1,
    ).font = TITLE_FONT

    row_number += 1

    headers = [
        "Sensor Type",
        "Sensor Name",
        "Records",
        "First Timestamp",
        "Last Timestamp",
    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            row=row_number,
            column=col,
            value=header,
        )

    style_header_row(
        ws,
        row_number,
    )

    row_number += 1

    grouped: dict[int, list[dict[str, Any]]] = (
        defaultdict(list)
    )

    for record in records:

        grouped[
            record["sensor_type"]
        ].append(record)

    for stype in sorted(grouped):

        group = grouped[stype]

        valid_timestamps = [
            r["timestamp_ms"]
            for r in group
            if r["timestamp_ms"] > 0
        ]

        ws.append(
            [
                stype,
                sensor_name(
                    stype,
                    group[0]["sensor_name"],
                ),
                len(group),
                (
                    format_timestamp_ms(
                        min(valid_timestamps)
                    )
                    if valid_timestamps
                    else "--"
                ),
                (
                    format_timestamp_ms(
                        max(valid_timestamps)
                    )
                    if valid_timestamps
                    else "--"
                ),
            ]
        )

    freeze_header(
        ws,
        row_number,
    )

    auto_width(ws)


# ============================================================
# Sensor worksheet
# ============================================================

def create_sensor_sheet(
    wb: Workbook,
    sensor_type: int,
    records: list[dict[str, Any]],
) -> None:
    """Write one stable worksheet per sensor type, preserving v2 provenance.

    The worksheet title is based on the Android sensor type, not the vendor
    reported hardware name. Hardware names remain row-level data so devices
    from different vendors do not change the expected sheet name.
    """

    category_name = SENSOR_NAMES.get(
        sensor_type,
        f"Sensor Type {sensor_type}",
    )
    sheet_name = unique_sheet_name(wb, category_name)
    ws = wb.create_sheet(title=sheet_name)

    ws["A1"] = f"SensorMax Dataset: {category_name}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Android Sensor Type {sensor_type}; hardware name is recorded per row"
    ws["A2"].font = SUBTITLE_FONT

    # Relative Time is a derived display column. All source contract fields
    # remain present so that v2 provenance is not lost during XLSX export.
    headers = [
        "Schema Version",
        "Sensor Timestamp (ns)",
        "Timestamp (ms)",
        "Relative Time (s)",
        "Machine ID",
        "Measurement Point",
        "Source Kind",
        "Source ID",
        "Channel ID",
        "Sensor Type",
        "Sensor Hardware Name",
        "Unit",
        "Value 0",
        "Value 1",
        "Value 2",
        "Value 3",
        "Value 4",
        "Value 5",
    ]
    header_row = 4
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    style_header_row(ws, header_row)

    sorted_records = sorted(
        records,
        key=lambda record: record.get("timestamp_ms", 0.0),
    )
    start_ts = (
        sorted_records[0].get("timestamp_ms", 0.0)
        if sorted_records
        else 0.0
    )

    value_keys = ("v0", "v1", "v2", "v3", "v4", "v5")
    for record in sorted_records:
        timestamp_ms = record.get("timestamp_ms")
        if timestamp_ms is None:
            relative_sec = None
        elif start_ts:
            relative_sec = (timestamp_ms - start_ts) / 1000.0
        else:
            relative_sec = 0.0

        sensor_timestamp_ns = record.get("sensor_timestamp_ns")
        # Excel stores only 15 significant decimal digits in numeric cells.
        # Keep nanosecond timestamps as text so a 16+ digit clock value remains exact.
        sensor_timestamp_ns_text = (
            str(sensor_timestamp_ns)
            if sensor_timestamp_ns is not None
            else None
        )
        values = [record.get(key) for key in value_keys]
        ws.append(
            [
                record.get("schema_version", 1),
                sensor_timestamp_ns_text,
                timestamp_ms,
                relative_sec,
                record.get("machine_id", ""),
                record.get("point", ""),
                record.get("source_kind", ""),
                record.get("source_id", ""),
                record.get("channel_id", ""),
                record.get("sensor_type", sensor_type),
                record.get("sensor_name", ""),
                record.get("unit", ""),
                *values,
            ]
        )

    first_data_row = header_row + 1
    last_data_row = header_row + len(sorted_records)
    for row in ws.iter_rows(
        min_row=first_data_row,
        max_row=last_data_row,
        min_col=1,
        max_col=len(headers),
    ):
        row[0].number_format = "0"  # schema version
        row[1].number_format = "@"  # exact sensor timestamp text
        row[2].number_format = "0"  # epoch milliseconds
        row[3].number_format = "0.0000"
        for cell in row[12:18]:
            cell.number_format = "0.000000"

    # Summary section only aggregates finite, available values. None values
    # remain blank, and measured numeric zero remains a valid zero.
    summary_row = header_row + len(sorted_records) + 3
    ws.cell(row=summary_row, column=1, value="Summary Statistics")
    ws.cell(row=summary_row, column=1).font = TITLE_FONT
    summary_row += 1

    summary_headers = ["Channel", "Minimum", "Maximum", "Mean", "RMS", "Unit"]
    for col, header in enumerate(summary_headers, start=1):
        ws.cell(row=summary_row, column=col, value=header)
    style_header_row(ws, summary_row)

    summary_data_start = summary_row + 1
    unit_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in sorted_records:
        reported_unit = clean_text(record.get("unit")) or "(unspecified)"
        unit_groups[reported_unit].append(record)

    # Never aggregate numerically incompatible units together. This matters
    # especially for external/custom sources that share the -1 sensor-type sentinel.
    for unit, unit_records in sorted(unit_groups.items()):
        stats = statistics(unit_records)
        for channel, key in (("Value 0", "v0"), ("Value 1", "v1"), ("Value 2", "v2")):
            ws.append([
                channel,
                stats[f"min_{key}"],
                stats[f"max_{key}"],
                stats[f"mean_{key}"],
                stats[f"rms_{key}"],
                unit,
            ])

    summary_data_end = ws.max_row
    for row in ws.iter_rows(
        min_row=summary_data_start,
        max_row=summary_data_end,
        min_col=2,
        max_col=5,
    ):
        for cell in row:
            cell.number_format = "0.000000"

    freeze_header(ws, header_row)
    if sorted_records:
        ws.auto_filter.ref = f"A{header_row}:R{last_data_row}"
    auto_width(ws)


# ============================================================
# Main conversion
# ============================================================

def convert_csv_to_excel(
    csv_path: Path,
    xlsx_path: Path,
) -> None:

    if not csv_path.exists():

        raise FileNotFoundError(
            f"CSV file not found: {csv_path}"
        )

    print(
        f"Reading: {csv_path}"
    )

    records, current_schema = read_csv(
        csv_path
    )

    if not records:

        print(
            "WARNING: No usable records found."
        )

        # Still produce a workbook so that
        # the pipeline has an explicit artifact.
        wb = Workbook()

        default_sheet = wb.active

        default_sheet.title = (
            "Session Metadata"
        )

        default_sheet["A1"] = (
            "SensorMax Dataset"
        )

        default_sheet["A1"].font = (
            TITLE_FONT
        )

        default_sheet["A3"] = (
            "No usable records found."
        )

        default_sheet["A4"] = (
            "Source"
        )

        default_sheet["B4"] = (
            csv_path.name
        )

        force_formula_like_input_to_text(wb)
        wb.save(xlsx_path)

        print(
            f"Generated empty workbook: {xlsx_path}"
        )

        return

    print(
        f"Records loaded: {len(records)}"
    )

    print(
        "Input schema: " +
        (
            "current Android SensorMax"
            if current_schema
            else "legacy compatibility"
        )
    )

    # --------------------------------------------------------
    # Group by sensor type
    # --------------------------------------------------------

    grouped: dict[
        int,
        list[dict[str, Any]]
    ] = defaultdict(list)

    for record in records:

        grouped[
            record["sensor_type"]
        ].append(record)

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    wb = Workbook()

    # Remove default sheet.
    default_sheet = wb.active
    wb.remove(default_sheet)

    # --------------------------------------------------------
    # Metadata
    # --------------------------------------------------------

    create_metadata_sheet(
        wb,
        csv_path,
        records,
        current_schema,
    )

    # --------------------------------------------------------
    # Sensor worksheets
    # --------------------------------------------------------

    for sensor_type in sorted(
        grouped.keys()
    ):

        create_sensor_sheet(
            wb,
            sensor_type,
            grouped[sensor_type],
        )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    xlsx_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    force_formula_like_input_to_text(wb)
    wb.save(
        xlsx_path
    )

    print(
        "Successfully generated:"
    )

    print(
        f"  {xlsx_path}"
    )

    print(
        f"  Sensor groups: {len(grouped)}"
    )

    for sensor_type in sorted(
        grouped
    ):

        print(
            f"    {sensor_type}: "
            f"{sensor_name(sensor_type)} "
            f"({len(grouped[sensor_type])} records)"
        )


# ============================================================
# CLI
# ============================================================

def build_parser() -> argparse.ArgumentParser:

    parser = argparse.ArgumentParser(
        description=(
            "Convert SensorMax raw sensor CSV "
            "to a multi-sheet Excel workbook."
        )
    )

    parser.add_argument(
        "csv_file",
        help=(
            "Input SensorMax raw CSV file"
        ),
    )

    parser.add_argument(
        "output_xlsx",
        nargs="?",
        help=(
            "Output Excel path. "
            "Defaults to <input>_processed.xlsx"
        ),
    )

    return parser


def main() -> int:

    parser = build_parser()

    args = parser.parse_args()

    csv_path = Path(
        args.csv_file
    ).expanduser().resolve()

    if args.output_xlsx:

        xlsx_path = Path(
            args.output_xlsx
        ).expanduser().resolve()

    else:

        xlsx_path = (
            csv_path.with_name(
                csv_path.stem +
                "_processed.xlsx"
            )
        )

    try:

        convert_csv_to_excel(
            csv_path,
            xlsx_path,
        )

        return 0

    except FileNotFoundError as exc:

        print(
            f"ERROR: {exc}"
        )

        return 1

    except PermissionError as exc:

        print(
            f"ERROR: Permission denied: {exc}"
        )

        return 1

    except Exception as exc:

        print(
            f"ERROR: Conversion failed: "
            f"{type(exc).__name__}: {exc}"
        )

        return 1


if __name__ == "__main__":
    raise SystemExit(
        main()
    )