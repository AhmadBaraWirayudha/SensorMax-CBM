"""Sprint 3 regression tests for the Android/network raw-measurement contract v2.

Run from the repository root:
    python -m unittest LaptopSuite.test_sprint3_measurement_contract -v
"""
from __future__ import annotations

import asyncio
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd

from LaptopSuite import export_converter, export_json, web_server


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ANDROID_ROOT = PROJECT_ROOT / "AndroidApp"


class OnePacketWebSocket:
    """Minimal async websocket fixture that sends one JSON message then closes."""

    remote_address = ("127.0.0.1", 51000)

    def __init__(self, packet: dict):
        self._messages = [json.dumps(packet, allow_nan=False)]
        self.request_headers = {}
        self.closed = False
        self.close_code = None
        self.close_reason = None

    def __aiter__(self):
        return self

    async def __anext__(self):
        if self._messages:
            return self._messages.pop(0)
        raise StopAsyncIteration

    async def close(self, code=None, reason=None):
        self.closed = True
        self.close_code = code
        self.close_reason = reason


class Sprint3MeasurementContractTests(unittest.TestCase):
    def test_contract_v2_header_matches_android_engine_source(self):
        engine = (
            ANDROID_ROOT / "app/src/main/java/com/research/sensormax/SensorEngine.kt"
        ).read_text(encoding="utf-8")
        expected_header = ",".join(web_server.RAW_CSV_COLUMNS)
        self.assertIn(f'"{expected_header}\\n"', engine)
        self.assertEqual(len(web_server.RAW_CSV_COLUMNS), 17)
        self.assertNotIn("Sensor_ID", web_server.RAW_CSV_COLUMNS)
        self.assertIn("Sensor_Type", web_server.RAW_CSV_COLUMNS)
        self.assertIn("Sensor_Name", web_server.RAW_CSV_COLUMNS)

    def test_android_stream_packet_includes_timestamp_source_and_unit_contract(self):
        stream_client = (
            ANDROID_ROOT / "app/src/main/java/com/research/sensormax/StreamClient.kt"
        ).read_text(encoding="utf-8")
        engine = (
            ANDROID_ROOT / "app/src/main/java/com/research/sensormax/SensorEngine.kt"
        ).read_text(encoding="utf-8")
        for field in (
            '"schemaVersion\\":2',
            '"sensorTimestampNs\\":',
            '"machineId\\":',
            '"sourceKind\\":',
            '"sourceId\\":',
            '"channelId\\":',
            '"sensorType\\":',
            '"unit\\":',
        ):
            self.assertIn(field, stream_client)
        self.assertIn("sensorTimestampNs = sensorTsNs", engine)
        self.assertIn("sensorType = sample.androidSensorType", engine)
        self.assertIn("unit = sample.unit", engine)
        self.assertIn("vals = sample.values", engine)
        self.assertIn("private fun csvField(value: Any?)", engine)

    def test_raw_packet_to_csv_row_preserves_identity_and_nulls(self):
        packet = {
            "kind": "raw",
            "schemaVersion": 2,
            "sensorTimestampNs": "9223372036854775000",
            "ts": 1700000000123,
            "machineId": "asset, west",
            "point": "DE, bearing \"A\"",
            "sourceKind": "ANDROID_SENSOR",
            "sourceId": "android, sensor",
            "channelId": "accelerometer|main",
            "sensorType": 1,
            "id": 1,
            "sensor": "Accelerometer, Rev \"B\"",
            "unit": "m/s²",
            "v0": 1.25,
            "v1": None,
            "v2": 0.0,
            "v3": None,
            "v4": None,
            "v5": None,
        }
        row = next(csv.reader(io.StringIO(web_server.raw_packet_to_csv_row(packet))))
        self.assertEqual(len(row), len(web_server.RAW_CSV_COLUMNS))
        record = dict(zip(web_server.RAW_CSV_COLUMNS, row))
        self.assertEqual(record["Schema_Version"], "2")
        self.assertEqual(record["Sensor_Timestamp_ns"], "9223372036854775000")
        self.assertEqual(record["Timestamp_ms"], "1700000000123")
        self.assertEqual(record["Machine_ID"], "asset, west")
        self.assertEqual(record["Point"], 'DE, bearing "A"')
        self.assertEqual(record["Source_ID"], "android, sensor")
        self.assertEqual(record["Sensor_Type"], "1")
        self.assertEqual(record["Sensor_Name"], 'Accelerometer, Rev "B"')
        self.assertEqual(record["Unit"], "m/s²")
        self.assertEqual(record["Val_0"], "1.25")
        self.assertEqual(record["Val_1"], "")
        self.assertEqual(record["Val_2"], "0.0")
        self.assertEqual(record["Val_3"], "")
        self.assertNotIn("Sensor_ID", record)

    def test_gateway_preserves_signed_64_bit_nanosecond_timestamps_exactly(self):
        timestamp_ns = "9223372036854775000"
        packet = {
            "kind": "raw", "schemaVersion": 2,
            "sensorTimestampNs": timestamp_ns, "ts": 1700000000123,
            "sensorType": 1, "sensor": "Accelerometer", "unit": "m/s²",
            "v0": 0.0, "v1": None, "v2": 2.0,
        }
        row = next(csv.reader(io.StringIO(web_server.raw_packet_to_csv_row(packet))))
        record = dict(zip(web_server.RAW_CSV_COLUMNS, row))
        self.assertEqual(record["Sensor_Timestamp_ns"], timestamp_ns)
        self.assertEqual(record["Val_0"], "0.0")
        self.assertEqual(record["Val_1"], "")

    def test_legacy_packet_is_normalized_without_fabricating_missing_axes(self):
        packet = {"ts": 1700000000000, "id": 1, "name": "Accelerometer", "x": 1.0, "z": 3.0}
        row = next(csv.reader(io.StringIO(web_server.raw_packet_to_csv_row(packet))))
        record = dict(zip(web_server.RAW_CSV_COLUMNS, row))
        self.assertEqual(record["Sensor_Type"], "1")
        self.assertEqual(record["Sensor_Name"], "Accelerometer")
        self.assertEqual(record["Val_0"], "1.0")
        self.assertEqual(record["Val_1"], "")
        self.assertEqual(record["Val_2"], "3.0")
        self.assertEqual(record["Sensor_Timestamp_ns"], "")
        self.assertEqual(record["Machine_ID"], "")

        non_android_packet = {
            "kind": "raw", "schemaVersion": 2, "sensorType": None, "id": -1,
            "sourceKind": "BLE_GATT", "sourceId": "ble-1", "channelId": "accel-x",
            "sensor": "External accelerometer", "ts": 500, "v0": 0.0,
        }
        external_row = next(csv.reader(io.StringIO(web_server.raw_packet_to_csv_row(non_android_packet))))
        external_record = dict(zip(web_server.RAW_CSV_COLUMNS, external_row))
        self.assertEqual(external_record["Sensor_Type"], "-1")
        self.assertEqual(external_record["Source_Kind"], "BLE_GATT")
        self.assertEqual(external_record["Val_0"], "0.0")

    def test_gateway_writes_exact_v2_header_without_blank_record(self):
        packet = {
            "kind": "raw", "schemaVersion": 2, "sensorTimestampNs": 123,
            "ts": 456, "machineId": "M,1", "point": "P1",
            "sourceKind": "ANDROID_SENSOR", "sourceId": "S1",
            "channelId": "C1", "sensorType": 1, "sensor": "Accelerometer",
            "unit": "m/s²", "v0": 9.81, "v1": None, "v2": 0.0,
            "v3": None, "v4": None, "v5": None,
        }
        async def exercise(log_dir: Path):
            prior = set(web_server.connected_clients)
            client = OnePacketWebSocket(packet)
            try:
                web_server.connected_clients.clear()
                with patch.object(web_server, "LOG_DIR", log_dir):
                    await web_server.handle_sensor_stream(client)
            finally:
                web_server.connected_clients.clear()
                web_server.connected_clients.update(prior)

        with tempfile.TemporaryDirectory() as td:
            log_dir = Path(td)
            asyncio.run(exercise(log_dir))
            outputs = list(log_dir.glob("*_raw.csv"))
            self.assertEqual(len(outputs), 1)
            with outputs[0].open("r", encoding="utf-8", newline="") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 1, "Header initialization must not create an empty row")
            self.assertEqual(tuple(rows[0].keys()), web_server.RAW_CSV_COLUMNS)
            self.assertEqual(rows[0]["Machine_ID"], "M,1")
            self.assertEqual(rows[0]["Val_1"], "")
            self.assertEqual(rows[0]["Val_2"], "0.0")

    def test_analysis_snapshot_is_separate_and_gateway_maps_its_header(self):
        packet = {
            "kind": "analysis", "schemaVersion": 1, "ts": 1700000000000,
            "machineId": "M1", "point": "DE-BRG-X", "sampleRateHz": 100.0,
            "overallRmsAccelMs2": 1.25, "overallRmsVelocityMmS": None,
            "overallRmsDisplacementUm": 4.5, "dominantAxis": "X",
            "dominantFreqHz": 30.0, "envelopePeakFreqHz": 60.0,
            "isoZone": "A", "bearingMatch": "NO_MATCH",
        }

        async def exercise(log_dir: Path):
            previous = set(web_server.connected_clients)
            client = OnePacketWebSocket(packet)
            try:
                web_server.connected_clients.clear()
                with patch.object(web_server, "LOG_DIR", log_dir):
                    await web_server.handle_sensor_stream(client)
            finally:
                web_server.connected_clients.clear()
                web_server.connected_clients.update(previous)

        with tempfile.TemporaryDirectory() as td:
            log_dir = Path(td)
            asyncio.run(exercise(log_dir))
            outputs = list(log_dir.glob("*_analysis.csv"))
            self.assertEqual(len(outputs), 1)
            with outputs[0].open("r", encoding="utf-8", newline="") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 1)
            self.assertEqual(tuple(rows[0].keys()), web_server.ANALYSIS_SNAPSHOT_CSV_COLUMNS)
            self.assertEqual(rows[0]["Timestamp_ms"], "1700000000000")
            self.assertEqual(rows[0]["Machine_ID"], "M1")
            self.assertEqual(rows[0]["RMS_Accel_ms2"], "1.25")
            self.assertEqual(rows[0]["RMS_Velocity_mms"], "")
            self.assertIn("val schemaVersion: Int = 1", (
                ANDROID_ROOT / "app/src/main/java/com/research/sensormax/StreamClient.kt"
            ).read_text(encoding="utf-8"))

    def test_converter_preserves_v2_provenance_and_missing_values(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "raw_v2.csv"
            with path.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=web_server.RAW_CSV_COLUMNS)
                writer.writeheader()
                writer.writerow({
                    "Schema_Version": 2,
                    "Sensor_Timestamp_ns": 9000000000,
                    "Timestamp_ms": 1700000000000,
                    "Machine_ID": "machine-A",
                    "Point": "DE-BRG-X",
                    "Source_Kind": "ANDROID_SENSOR",
                    "Source_ID": "phone-accelerometer",
                    "Channel_ID": "accelerometer-main",
                    "Sensor_Type": 1,
                    "Sensor_Name": "Accelerometer",
                    "Unit": "m/s²",
                    "Val_0": "1.25",
                    "Val_1": "",
                    "Val_2": "0",
                    "Val_3": "NaN",
                    "Val_4": "",
                    "Val_5": "",
                })
            records, is_current = export_converter.read_csv(path)
            self.assertTrue(is_current)
            self.assertEqual(len(records), 1)
            record = records[0]
            self.assertEqual(record["schema_version"], 2)
            self.assertEqual(record["sensor_timestamp_ns"], 9000000000)
            self.assertEqual(record["source_id"], "phone-accelerometer")
            self.assertEqual(record["channel_id"], "accelerometer-main")
            self.assertEqual(record["unit"], "m/s²")
            self.assertEqual(record["v0"], 1.25)
            self.assertIsNone(record["v1"])
            self.assertEqual(record["v2"], 0.0)
            self.assertIsNone(record["v3"])
            stats = export_converter.statistics(records)
            self.assertEqual(stats["min_v0"], 1.25)
            self.assertIsNone(stats["min_v1"], "Missing channels must not contribute a fake zero")

    def test_xlsx_keeps_missing_values_blank_in_cells_and_statistics(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "raw_v2.csv"
            output = Path(td) / "raw_v2.xlsx"
            with source.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=web_server.RAW_CSV_COLUMNS)
                writer.writeheader()
                writer.writerow({
                    "Schema_Version": 2,
                    "Sensor_Timestamp_ns": 123,
                    "Timestamp_ms": 1700000000000,
                    "Machine_ID": "M1",
                    "Point": "P1",
                    "Source_Kind": "ANDROID_SENSOR",
                    "Source_ID": "accel-1",
                    "Channel_ID": "accelerometer",
                    "Sensor_Type": 1,
                    "Sensor_Name": "Accelerometer",
                    "Unit": "m/s²",
                    "Val_0": "1.5",
                    "Val_1": "",
                    "Val_2": "0",
                    "Val_3": "",
                    "Val_4": "",
                    "Val_5": "",
                })
            export_converter.convert_csv_to_excel(source, output)
            workbook = openpyxl.load_workbook(output, data_only=True, read_only=False)
            try:
                sheet = workbook["Accelerometer"]
                header = {sheet.cell(4, col).value: col for col in range(1, sheet.max_column + 1)}
                self.assertEqual(sheet.cell(5, header["Value 0"]).value, 1.5)
                self.assertIsNone(sheet.cell(5, header["Value 1"]).value)
                self.assertEqual(sheet.cell(5, header["Value 2"]).value, 0)
                summary_title_row = 4 + 1 + 3
                summary_headers = {
                    sheet.cell(summary_title_row + 1, col).value: col
                    for col in range(1, sheet.max_column + 1)
                }
                summary_values = {}
                for offset in range(2, 5):
                    row_number = summary_title_row + offset
                    channel = sheet.cell(row_number, summary_headers["Channel"]).value
                    summary_values[channel] = {
                        "min": sheet.cell(row_number, summary_headers["Minimum"]).value,
                        "rms": sheet.cell(row_number, summary_headers["RMS"]).value,
                    }
                self.assertIsNone(summary_values["Value 1"]["min"])
                self.assertIsNone(summary_values["Value 1"]["rms"])
                self.assertEqual(summary_values["Value 2"]["min"], 0)
                self.assertEqual(sheet.auto_filter.ref, "A4:R5")
            finally:
                workbook.close()

    def test_xlsx_uses_stable_sensor_type_worksheet_names_and_preserves_v2_provenance(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "multi_sensor_v2.csv"
            output = Path(td) / "multi_sensor_v2.xlsx"
            rows = [
                {
                    "Schema_Version": 2, "Sensor_Timestamp_ns": "1700000000123456789",
                    "Timestamp_ms": 1700000000000, "Machine_ID": "machine-A", "Point": "DE-BRG-X",
                    "Source_Kind": "ANDROID_SENSOR", "Source_ID": "phone-A", "Channel_ID": "accel-main",
                    "Sensor_Type": 1, "Sensor_Name": "Bosch BMI160 Accelerometer", "Unit": "m/s²",
                    "Val_0": "1.25", "Val_1": "", "Val_2": "0", "Val_3": "NaN", "Val_4": "", "Val_5": "",
                },
                {
                    "Schema_Version": 2, "Sensor_Timestamp_ns": "1700000001123456789",
                    "Timestamp_ms": 1700000001000, "Machine_ID": "machine-A", "Point": "DE-BRG-X",
                    "Source_Kind": "ANDROID_SENSOR", "Source_ID": "phone-A", "Channel_ID": "accel-main",
                    "Sensor_Type": 1, "Sensor_Name": "Bosch BMI160 Accelerometer", "Unit": "m/s²",
                    "Val_0": "2.5", "Val_1": "", "Val_2": "0", "Val_3": "", "Val_4": "", "Val_5": "",
                },
                {
                    "Schema_Version": 2, "Sensor_Timestamp_ns": "1700000002123456789",
                    "Timestamp_ms": 1700000002000, "Machine_ID": "machine-A", "Point": "DE-BRG-X",
                    "Source_Kind": "ANDROID_SENSOR", "Source_ID": "phone-A", "Channel_ID": "gyro-main",
                    "Sensor_Type": 4, "Sensor_Name": "ICM42688 Gyroscope", "Unit": "rad/s",
                    "Val_0": "0.1", "Val_1": "0.2", "Val_2": "0.3", "Val_3": "", "Val_4": "", "Val_5": "",
                },
            ]
            with source.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=web_server.RAW_CSV_COLUMNS)
                writer.writeheader()
                writer.writerows(rows)

            export_converter.convert_csv_to_excel(source, output)
            workbook = openpyxl.load_workbook(output, data_only=True, read_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Session Metadata", "Accelerometer", "Gyroscope"])
                accel = workbook["Accelerometer"]
                self.assertEqual(accel.max_row, 13)
                header = {accel.cell(4, col).value: col for col in range(1, 19)}
                self.assertEqual(set(header), {
                    "Schema Version", "Sensor Timestamp (ns)", "Timestamp (ms)", "Relative Time (s)",
                    "Machine ID", "Measurement Point", "Source Kind", "Source ID", "Channel ID",
                    "Sensor Type", "Sensor Hardware Name", "Unit", "Value 0", "Value 1", "Value 2",
                    "Value 3", "Value 4", "Value 5",
                })
                # Two data records at rows 5-6, plus stable filter range.
                self.assertEqual(accel.auto_filter.ref, "A4:R6")
                self.assertEqual(accel.cell(5, header["Schema Version"]).value, 2)
                # Nanosecond timestamps are stored as text to avoid Excel's 15-digit precision loss.
                self.assertEqual(accel.cell(5, header["Sensor Timestamp (ns)"]).value, "1700000000123456789")
                self.assertEqual(accel.cell(5, header["Timestamp (ms)"]).value, 1700000000000)
                self.assertEqual(accel.cell(5, header["Machine ID"]).value, "machine-A")
                self.assertEqual(accel.cell(5, header["Measurement Point"]).value, "DE-BRG-X")
                self.assertEqual(accel.cell(5, header["Source Kind"]).value, "ANDROID_SENSOR")
                self.assertEqual(accel.cell(5, header["Source ID"]).value, "phone-A")
                self.assertEqual(accel.cell(5, header["Channel ID"]).value, "accel-main")
                self.assertEqual(accel.cell(5, header["Sensor Type"]).value, 1)
                self.assertEqual(accel.cell(5, header["Sensor Hardware Name"]).value, "Bosch BMI160 Accelerometer")
                self.assertEqual(accel.cell(5, header["Unit"]).value, "m/s²")
                self.assertEqual(accel.cell(5, header["Value 0"]).value, 1.25)
                self.assertIsNone(accel.cell(5, header["Value 1"]).value)
                self.assertEqual(accel.cell(5, header["Value 2"]).value, 0)
                self.assertIsNone(accel.cell(5, header["Value 3"]).value)

                gyro = workbook["Gyroscope"]
                gyro_header = {gyro.cell(4, col).value: col for col in range(1, 19)}
                self.assertEqual(gyro.auto_filter.ref, "A4:R5")
                self.assertEqual(gyro.cell(5, gyro_header["Sensor Hardware Name"]).value, "ICM42688 Gyroscope")
                self.assertEqual(gyro.cell(5, gyro_header["Unit"]).value, "rad/s")

                metadata = workbook["Session Metadata"]
                self.assertEqual(metadata["B6"].value, 3)
            finally:
                workbook.close()

    def test_xlsx_legacy_schema_retains_compatibility_and_leaves_v2_fields_blank(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "legacy_v1.csv"
            output = Path(td) / "legacy_v1.xlsx"
            columns = ["Timestamp_ms", "Machine_ID", "Point", "Sensor_Type", "Sensor_Name", "Val_0", "Val_1", "Val_2", "Val_3", "Val_4", "Val_5"]
            with source.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=columns)
                writer.writeheader()
                writer.writerow({
                    "Timestamp_ms": 1700000000000, "Machine_ID": "legacy-M", "Point": "legacy-P",
                    "Sensor_Type": 1, "Sensor_Name": "Legacy accelerometer", "Val_0": "0",
                    "Val_1": "", "Val_2": "3.0", "Val_3": "", "Val_4": "", "Val_5": "",
                })
            export_converter.convert_csv_to_excel(source, output)
            workbook = openpyxl.load_workbook(output, data_only=True, read_only=True)
            try:
                self.assertIn("Accelerometer", workbook.sheetnames)
                sheet = workbook["Accelerometer"]
                header = {sheet.cell(4, col).value: col for col in range(1, 19)}
                self.assertEqual(sheet.cell(5, header["Schema Version"]).value, 1)
                self.assertIsNone(sheet.cell(5, header["Sensor Timestamp (ns)"]).value)
                self.assertEqual(sheet.cell(5, header["Sensor Hardware Name"]).value, "Legacy accelerometer")
                self.assertIsNone(sheet.cell(5, header["Source ID"]).value)
                self.assertEqual(sheet.cell(5, header["Value 0"]).value, 0)
                self.assertIsNone(sheet.cell(5, header["Value 1"]).value)
                self.assertEqual(sheet.cell(5, header["Value 2"]).value, 3)
            finally:
                workbook.close()


    def test_xlsx_summary_does_not_mix_units_for_external_sensor_type(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "mixed_units.csv"
            output = Path(td) / "mixed_units.xlsx"
            rows = [
                {
                    "Schema_Version": 2, "Sensor_Timestamp_ns": 101, "Timestamp_ms": 1700000000000,
                    "Machine_ID": "M1", "Point": "P1", "Source_Kind": "BLE_GATT", "Source_ID": "ble-a",
                    "Channel_ID": "accel", "Sensor_Type": -1, "Sensor_Name": "External acceleration",
                    "Unit": "m/s²", "Val_0": "1", "Val_1": "", "Val_2": "", "Val_3": "", "Val_4": "", "Val_5": "",
                },
                {
                    "Schema_Version": 2, "Sensor_Timestamp_ns": 102, "Timestamp_ms": 1700000001000,
                    "Machine_ID": "M1", "Point": "P1", "Source_Kind": "BLE_GATT", "Source_ID": "ble-b",
                    "Channel_ID": "angular-rate", "Sensor_Type": -1, "Sensor_Name": "External gyro",
                    "Unit": "rad/s", "Val_0": "100", "Val_1": "", "Val_2": "", "Val_3": "", "Val_4": "", "Val_5": "",
                },
            ]
            with source.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=web_server.RAW_CSV_COLUMNS)
                writer.writeheader()
                writer.writerows(rows)
            export_converter.convert_csv_to_excel(source, output)
            workbook = openpyxl.load_workbook(output, data_only=True, read_only=True)
            try:
                sheet = workbook["Sensor Type -1"]
                summary_title_row = 4 + len(rows) + 3
                summary_headers = {
                    sheet.cell(summary_title_row + 1, col).value: col
                    for col in range(1, 7)
                }
                summary = {}
                for row_num in range(summary_title_row + 2, summary_title_row + 8):
                    channel = sheet.cell(row_num, summary_headers["Channel"]).value
                    unit = sheet.cell(row_num, summary_headers["Unit"]).value
                    summary[(channel, unit)] = sheet.cell(row_num, summary_headers["Minimum"]).value
                self.assertEqual(summary[("Value 0", "m/s²")], 1)
                self.assertEqual(summary[("Value 0", "rad/s")], 100)
            finally:
                workbook.close()

    def test_xlsx_treats_formula_like_metadata_as_literal_text(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "formula_text.csv"
            output = Path(td) / "formula_text.xlsx"
            formula_like = '=HYPERLINK("https://example.test","click")'
            columns = list(web_server.RAW_CSV_COLUMNS)
            row = {
                "Schema_Version": 2, "Sensor_Timestamp_ns": "1234567890123456789",
                "Timestamp_ms": 1700000000000, "Machine_ID": "M1", "Point": "P1",
                "Source_Kind": "BLE_GATT", "Source_ID": formula_like, "Channel_ID": "C1",
                "Sensor_Type": 1, "Sensor_Name": "=1+1", "Unit": "m/s²",
                "Val_0": "0", "Val_1": "", "Val_2": "", "Val_3": "", "Val_4": "", "Val_5": "",
            }
            with source.open("w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=columns)
                writer.writeheader()
                writer.writerow(row)
            export_converter.convert_csv_to_excel(source, output)
            workbook = openpyxl.load_workbook(output, data_only=False, read_only=True)
            try:
                sheet = workbook["Accelerometer"]
                headers = {sheet.cell(4, col).value: col for col in range(1, 19)}
                source_cell = sheet.cell(5, headers["Source ID"])
                name_cell = sheet.cell(5, headers["Sensor Hardware Name"])
                self.assertEqual(source_cell.value, formula_like)
                self.assertEqual(source_cell.data_type, "s")
                self.assertEqual(name_cell.value, "=1+1")
                self.assertEqual(name_cell.data_type, "s")
            finally:
                workbook.close()

    def test_android_telemetry_does_not_invent_channels_for_short_samples(self):
        root = Path(__file__).resolve().parents[1]
        engine = (root / "AndroidApp/app/src/main/java/com/research/sensormax/SensorEngine.kt").read_text(encoding="utf-8")
        telemetry = (root / "AndroidApp/app/src/main/java/com/research/sensormax/SensorTelemetry.kt").read_text(encoding="utf-8")
        self.assertIn("val telemetryValues = sample.values.copyOf()", engine)
        self.assertNotIn("FloatArray(6) { index ->", engine)
        self.assertIn('value.takeIf { it.isFinite() }', engine)
        self.assertIn('else "—"', telemetry)
        self.assertIn('values.take(3).all { it.isFinite() }', telemetry)

    def test_json_export_distinguishes_raw_v2_from_legacy_v1(self):
        columns_v2 = list(web_server.RAW_CSV_COLUMNS)
        columns_v1 = [
            "Timestamp_ms", "Machine_ID", "Point", "Sensor_Type", "Sensor_Name",
            "Val_0", "Val_1", "Val_2", "Val_3", "Val_4", "Val_5",
        ]
        self.assertEqual(export_json.detect_schema(pd.DataFrame(columns=columns_v2)), "android_raw_v2")
        self.assertEqual(export_json.detect_schema(pd.DataFrame(columns=columns_v1)), "android_raw_v1")

    def test_json_export_keeps_nulls_and_extra_v2_columns(self):
        frame = pd.DataFrame([{
            "Schema_Version": 2,
            "Sensor_Timestamp_ns": 123,
            "Timestamp_ms": 456,
            "Machine_ID": "M1",
            "Point": "P1",
            "Source_Kind": "BLE_GATT",
            "Source_ID": "ble-device-1",
            "Channel_ID": "accel-x",
            "Sensor_Type": -1,
            "Sensor_Name": "External accelerometer",
            "Unit": "m/s²",
            "Val_0": 0.0,
            "Val_1": float("nan"),
            "Val_2": 1.0,
            "Val_3": float("nan"),
            "Val_4": float("nan"),
            "Val_5": float("nan"),
        }])
        telemetry = export_json.build_raw_telemetry(frame)
        self.assertEqual(telemetry[0]["Source_ID"], "ble-device-1")
        self.assertEqual(telemetry[0]["Sensor_Timestamp_ns"], "123")
        self.assertIsNone(telemetry[0]["Val_1"])
        self.assertIsNone(telemetry[0]["Val_3"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
